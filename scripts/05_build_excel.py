"""Step 5 - Build the Excel workbook: order sample, lookup sheet and a live-formula KPI sheet."""
import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from config import EXCEL_FILE, PROCESSED, STATE_NAMES, make_dirs

SAMPLE_ROWS = 5000
COLS = ["order_id", "order_date", "order_month", "customer_state", "main_category", "items",
        "order_value", "freight_value", "payment_type", "delivery_days", "delivery_status", "review_score"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT


def main():
    path = PROCESSED / "orders_master.csv"
    if not path.exists():
        raise SystemExit("orders_master.csv not found. Run scripts/03_clean_eda.py first.")
    make_dirs()
    orders = pd.read_csv(path)
    orders = orders[orders["order_value"].notna()][COLS]
    sample = orders.sample(min(SAMPLE_ROWS, len(orders)), random_state=42).sort_values("order_date")
    last = len(sample) + 1
    states = pd.DataFrame(sorted(STATE_NAMES.items()), columns=["state_code", "state_name"])
    top_states = sample.groupby("customer_state")["order_value"].sum().nlargest(10).index.tolist()

    with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl") as xl:
        sample.to_excel(xl, sheet_name="Orders", index=False)
        states.to_excel(xl, sheet_name="State_Lookup", index=False)
        wb = xl.book

        ws = wb["Orders"]
        style_header(ws)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        # lookup column so the workbook shows VLOOKUP against the State_Lookup sheet
        ws["M1"] = "state_name"
        for r in range(2, last + 1):
            ws[f"M{r}"] = f"=VLOOKUP(D{r},State_Lookup!$A$2:$B$28,2,FALSE)"
        style_header(ws)
        for i in range(1, 14):
            ws.column_dimensions[get_column_letter(i)].width = 18
        style_header(wb["State_Lookup"])

        k = wb.create_sheet("KPI_Summary", 0)
        rng = lambda col: f"Orders!${col}$2:${col}${last}"
        k["A1"] = "KPI Summary (live formulas on the Orders sheet)"
        k["A1"].font = Font(bold=True, size=14)
        kpis = [
            ("Total orders", f"=COUNTA({rng('A')})", "#,##0"),
            ("Total revenue", f"=SUM({rng('G')})", "#,##0"),
            ("Average order value", f"=AVERAGE({rng('G')})", "#,##0.00"),
            ("Average delivery days", f"=AVERAGE({rng('J')})", "0.0"),
            ("Late delivery %", f'=COUNTIF({rng("K")},"Late")/COUNTIF({rng("K")},"<>Not Delivered")', "0.00%"),
            ("Average review score", f"=AVERAGE({rng('L')})", "0.00"),
            ("Avg review - on time", f'=AVERAGEIFS({rng("L")},{rng("K")},"On Time")', "0.00"),
            ("Avg review - late", f'=AVERAGEIFS({rng("L")},{rng("K")},"Late")', "0.00"),
        ]
        for i, (label, formula, fmt) in enumerate(kpis, start=3):
            k[f"A{i}"], k[f"B{i}"] = label, formula
            k[f"B{i}"].number_format = fmt

        head = 13
        for j, title in enumerate(["State", "State name", "Orders", "Revenue", "Avg order value",
                                   "Late %", "Avg review"], start=1):
            k.cell(row=head, column=j, value=title)
        style_header(k, head)
        for i, st in enumerate(top_states, start=head + 1):
            k[f"A{i}"] = st
            k[f"B{i}"] = f"=VLOOKUP(A{i},State_Lookup!$A$2:$B$28,2,FALSE)"
            k[f"C{i}"] = f"=COUNTIFS({rng('D')},A{i})"
            k[f"D{i}"] = f"=SUMIFS({rng('G')},{rng('D')},A{i})"
            k[f"E{i}"] = f"=IFERROR(D{i}/C{i},0)"
            k[f"F{i}"] = f'=IFERROR(COUNTIFS({rng("D")},A{i},{rng("K")},"Late")/C{i},0)'
            k[f"G{i}"] = f"=IFERROR(AVERAGEIFS({rng('L')},{rng('D')},A{i}),0)"
            k[f"D{i}"].number_format = k[f"E{i}"].number_format = "#,##0"
            k[f"F{i}"].number_format = "0.00%"
            k[f"G{i}"].number_format = "0.00"
        end = head + len(top_states)
        k.conditional_formatting.add(f"D{head + 1}:D{end}", DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
        k.conditional_formatting.add(f"F{head + 1}:F{end}", ColorScaleRule(
            start_type="min", start_color="63BE7B", end_type="max", end_color="F8696B"))
        k.conditional_formatting.add(f"G{head + 1}:G{end}", ColorScaleRule(
            start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
        for col, width in zip("ABCDEFG", [26, 22, 12, 16, 18, 12, 14]):
            k.column_dimensions[col].width = width

    print(f"Excel workbook ready: {EXCEL_FILE} ({len(sample):,} order rows)")


if __name__ == "__main__":
    main()
