"""Step 8 - Build a short stakeholder presentation (PowerPoint) from the outputs."""
from pptx import Presentation
from pptx.util import Inches, Pt

from config import FIGURES, OUTPUTS, ROOT

PPT_FILE = ROOT / "presentation" / "olist_insights_deck.pptx"

SLIDES = [
    ("Where does revenue come from?", "01_monthly_revenue.png",
     "Revenue grew through 2017 and levelled off in 2018. Three states (SP, RJ, MG) drive 63% of revenue."),
    ("Which categories matter most?", "02_top_categories.png",
     "Health & beauty, watches & gifts and bed, bath & table are the top three categories."),
    ("Do customers come back?", "06_cohort_retention.png",
     "Only 3% of customers order again. Retention after the first month is below 1%."),
    ("Does late delivery hurt reviews?", "05_review_by_delivery_status.png",
     "Late orders average 2.57 stars versus 4.29 for on-time orders (p < 0.001)."),
]
RECOMMENDATIONS = [
    "Fix late delivery first: set realistic estimates for northern states and review slow sellers.",
    "Launch a repeat-purchase offer in the first 30 days for high-value new customers.",
    "Win back the high-value 'At Risk' segment, which holds 29% of revenue.",
]


def bullets(slide, lines, top=1.6, size=18):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12), Inches(4.5)).text_frame
    box.word_wrap = True
    for i, line in enumerate(lines):
        p = box.paragraphs[0] if i == 0 else box.add_paragraph()
        p.text = line
        p.font.size = Pt(size)


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Olist E-Commerce: Sales, Customers and Delivery"
    s.placeholders[1].text = "Analysis of 99K orders using SQL, Python, Excel and Tableau"

    summary = (OUTPUTS / "insights_summary.txt").read_text(encoding="utf-8").splitlines()[2:11]
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Key numbers"
    bullets(s, summary, size=20)

    for title, image, takeaway in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        s.shapes.add_picture(str(FIGURES / image), Inches(1.6), Inches(1.4), height=Inches(4.6))
        bullets(s, [takeaway], top=6.2, size=18)

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Recommendations"
    bullets(s, [f"{i}. {r}" for i, r in enumerate(RECOMMENDATIONS, 1)], size=22)

    PPT_FILE.parent.mkdir(exist_ok=True)
    prs.save(PPT_FILE)
    print(f"Presentation ready: {PPT_FILE} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
